import docx
from xml.etree import ElementTree as ET

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

print("=== PAGE BREAKS IN CORRECT DOC ===")
for i, p in enumerate(doc.paragraphs):
    xml = ET.tostring(p._p, encoding='unicode')
    if 'w:br' in xml or 'page' in xml:
        print(f"P{i}: {xml}")

print("\n=== PAGE BREAKS IN OLD DOC ===")
doc_o = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")
for i, p in enumerate(doc_o.paragraphs):
    xml = ET.tostring(p._p, encoding='unicode')
    if 'w:br' in xml or 'page' in xml:
        print(f"P{i}: {xml}")

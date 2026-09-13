import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from xml.etree import ElementTree as ET
import os

TEMPLATE_PATH = "artifacts/Property Acknowledgement Document Template.docx"
GOLDEN_PATH = "docs/Property Acknowledgement Document Correct Format.docx"

def stabilize_table_borders(table: docx.table.Table) -> None:
    tbl_borders = parse_xml(
        r"""
        <w:tblBorders {}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>
        </w:tblBorders>
        """.format(nsdecls("w"))
    )
    # Remove existing tblBorders if present to avoid duplication
    existing = table._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders')
    if existing is not None:
        table._tbl.tblPr.remove(existing)
    table._tbl.tblPr.append(tbl_borders)

def set_cell_checkbox(cell, checked: bool):
    """Sets paragraph in cell to numId=999 (checked) or numId=998 (unchecked)."""
    p = cell.paragraphs[0] if cell.paragraphs else None
    if p is not None:
        numId_elem = p._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId')
        if numId_elem is not None:
            numId_elem.set(qn('w:val'), '999' if checked else '998')

def populate_doc(data, template_path=TEMPLATE_PATH):
    doc = docx.Document(template_path)

    # 1. Update Form No (Paragraph 5)
    form_no = str(data.get("form_no", "0257"))
    p5 = doc.paragraphs[5]
    # Remove any runs after "No: "
    no_run_idx = -1
    for idx, r in enumerate(p5.runs):
        if "No:" in r.text:
            no_run_idx = idx
            break
    if no_run_idx != -1:
        # Keep runs up to no_run_idx, remove following
        while len(p5.runs) > no_run_idx + 1:
            p5._p.remove(p5.runs[-1]._r)
        # Add runs for each digit in form_no
        for digit in form_no:
            r = p5.add_run(digit)
            r.bold = True
            r.font.size = Pt(14.0)
            r.font.color.rgb = RGBColor(255, 0, 0)

    # 2. Update Header Date (Paragraph 7)
    header_date = str(data.get("date", "27.03.2025"))
    p7 = doc.paragraphs[7]
    date_run_found = False
    for r in p7.runs:
        if r.text.strip() == "27.03.2025" or (r.text.strip() and r.text.strip()[0].isdigit() and "." in r.text):
            r.text = header_date
            r.bold = True
            date_run_found = True
            break
    if not date_run_found and len(p7.runs) > 12:
        p7.runs[12].text = header_date
        p7.runs[12].bold = True

    # 3. Table 0 (Customer Particulars & Requirements)
    t0 = doc.tables[0]
    stabilize_table_borders(t0)
    c0 = t0.rows[0].cells[0]
    c1 = t0.rows[0].cells[1]

    # Find nested tables in c0
    c0_nested_tbls = [docx.table.Table(node, c0) for node in c0._tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')]
    
    # Salutation in c0_nested_tbls[0]: 1 row x 6 cols
    salutation = (data.get("salutation") or "").upper()
    if c0_nested_tbls:
        t_sal = c0_nested_tbls[0]
        # Col 0: MR, Col 2: MRS, Col 4: MS
        set_cell_checkbox(t_sal.rows[0].cells[0], "MR" in salutation and "MRS" not in salutation)
        set_cell_checkbox(t_sal.rows[0].cells[2], "MRS" in salutation)
        set_cell_checkbox(t_sal.rows[0].cells[4], "MS" in salutation)

    # Customer Name: c0.paragraphs[0]
    cust_name = data.get("customer_name") or data.get("name") or ""
    p_name = c0.paragraphs[0]
    # Set text after "Full Name:  "
    name_run = None
    for r in p_name.runs:
        if r.text.strip() in ["Nick", "Lee Wan Soon"] or (name_run is None and r.text not in ["Full ", "Name:", " "]):
            name_run = r
    if name_run is not None:
        name_run.text = cust_name
    elif len(p_name.runs) > 4:
        p_name.runs[4].text = cust_name
    elif cust_name:
        p_name.add_run(cust_name)

    # No of pax: c0.paragraphs[2]
    pax = str(data.get("no_of_pax", "2"))
    p_pax = c0.paragraphs[2]
    for r in p_pax.runs:
        if r.text.strip().isdigit():
            r.text = pax
            r.bold = True
            break

    # Company Name: c0.paragraphs[4]
    company_name = data.get("company_name", "")
    reg_no = data.get("company_reg_no", "")
    company_full = f"{company_name} ({reg_no})" if reg_no else company_name
    p_comp = c0.paragraphs[4]
    if len(p_comp.runs) > 1:
        p_comp.runs[1].text = company_full
        if company_full:
            p_comp.runs[1].font.highlight_color = WD_COLOR_INDEX.YELLOW
            p_comp.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p_comp.runs[1].font.highlight_color = None
            p_comp.runs[0].font.highlight_color = None

    # Company Address: c0.paragraphs[6] and c0.paragraphs[7]
    company_addr = data.get("company_address", "")
    p_addr1 = c0.paragraphs[6]
    p_addr2 = c0.paragraphs[7]
    if company_addr:
        addr_lines = company_addr.split("\n")
        line1 = addr_lines[0].strip()
        line2 = "\n".join(addr_lines[1:]).strip() if len(addr_lines) > 1 else ""
        if len(p_addr1.runs) > 3:
            p_addr1.runs[3].text = line1
        if len(p_addr2.runs) > 0:
            p_addr2.runs[0].text = line2
    else:
        if len(p_addr1.runs) > 3:
            p_addr1.runs[3].text = ""
        if len(p_addr2.runs) > 0:
            p_addr2.runs[0].text = ""

    # Car plate no: c0.paragraphs[9]
    car_plate = data.get("car_plate", "")
    p_car = c0.paragraphs[9]
    if car_plate:
        if len(p_car.runs) > 1:
            p_car.runs[1].text = car_plate
        else:
            r = p_car.add_run(car_plate)
            r.bold = True
            r.font.size = Pt(10.0)

    # Tel (Hp): c0.paragraphs[10]
    phone = data.get("phone", "")
    p_tel = c0.paragraphs[10]
    if len(p_tel.runs) > 1:
        p_tel.runs[1].text = phone
    elif phone:
        p_tel.add_run(phone)

    # Called in date: c0.paragraphs[11]
    called_in = data.get("called_in_date", header_date)
    p_called = c0.paragraphs[11]
    if len(p_called.runs) > 2:
        p_called.runs[2].text = called_in
        p_called.runs[2].bold = True

    # Marketing Referral Sources in c0_nested_tbls[1]: 6 rows x 2 cols
    ref_source = (data.get("referral_source") or "").lower()
    if len(c0_nested_tbls) > 1:
        t_ref = c0_nested_tbls[1]
        source_keys = ["iproperty", "mudah", "bentongland", "banner", "whatsapp", "facebook"]
        for r_i, key in enumerate(source_keys):
            if r_i < len(t_ref.rows):
                is_match = key in ref_source
                set_cell_checkbox(t_ref.rows[r_i].cells[0], is_match)

    # 4. Table 0 Cell 1 (Customer Requirements)
    # Customer Request: c1.paragraphs[1]
    cust_req = data.get("customer_request", "")
    p_req = c1.paragraphs[1]
    if cust_req:
        while len(p_req.runs) > 1:
            p_req._p.remove(p_req.runs[-1]._r)
        r = p_req.add_run(cust_req)
        r.bold = True
        r.font.size = Pt(10.0)

    # Requirement summary: c1.paragraphs[5]
    req_summary = data.get("requirement_summary", "")
    p_sum = c1.paragraphs[5]
    if req_summary:
        if len(p_sum.runs) > 2:
            p_sum.runs[2].text = req_summary
        elif len(p_sum.runs) == 2:
            p_sum.add_run(req_summary)

    # Property Types in c1 nested table 0 (4 rows x 4 cols)
    c1_nested_tbls = [docx.table.Table(node, c1) for node in c1._tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')]
    prop_types = [pt.lower() for pt in data.get("property_types", [])]
    if c1_nested_tbls:
        t_props = c1_nested_tbls[0]
        # Row 0 Col 0: Agricultural Land
        set_cell_checkbox(t_props.rows[0].cells[0], any("agri" in pt for pt in prop_types))
        # Row 1 Col 0: Commercial Land
        set_cell_checkbox(t_props.rows[1].cells[0], any("comm" in pt for pt in prop_types))
        # Row 1 Col 2: Hotel / Resort
        set_cell_checkbox(t_props.rows[1].cells[2], any("hotel" in pt or "resort" in pt for pt in prop_types))
        # Row 2 Col 0: Industrial Land
        set_cell_checkbox(t_props.rows[2].cells[0], any("ind" in pt for pt in prop_types))
        # Row 2 Col 2: Factory
        set_cell_checkbox(t_props.rows[2].cells[2], any("factory" in pt for pt in prop_types))
        # Row 3 Col 0: Residential Land
        set_cell_checkbox(t_props.rows[3].cells[0], any("res" in pt for pt in prop_types))
        # Row 3 Col 2: Condominium
        set_cell_checkbox(t_props.rows[3].cells[2], any("condo" in pt for pt in prop_types))

    # Location: c1.paragraphs[9]
    loc = data.get("target_location", "")
    p_loc = c1.paragraphs[9]
    if loc:
        if len(p_loc.runs) > 2:
            p_loc.runs[2].text = loc
        elif len(p_loc.runs) == 2:
            p_loc.add_run(loc)

    # Remarks: c1.paragraphs[11]
    remarks = data.get("remarks", "")
    p_rem = c1.paragraphs[11]
    if remarks:
        while len(p_rem.runs) > 2:
            p_rem._p.remove(p_rem.runs[-1]._r)
        if len(p_rem.runs) == 2:
            p_rem.add_run(remarks)
        elif len(p_rem.runs) > 2:
            p_rem.runs[2].text = remarks

    # Assigned To in c1 nested table 1 (2 rows x 2 cols)
    assigned = (data.get("assigned_to") or "").lower()
    if len(c1_nested_tbls) > 1:
        t_ass = c1_nested_tbls[1]
        is_direct = "direct" in assigned
        is_cobroke = "co-broke" in assigned or "cobroke" in assigned
        if not is_direct and not is_cobroke:
            is_direct = True # default
        set_cell_checkbox(t_ass.rows[0].cells[0], is_cobroke)
        set_cell_checkbox(t_ass.rows[1].cells[0], is_direct)

    # Legal Partner Agency (Paragraph 9)
    agency_name = data.get("partner_agency")
    if agency_name and len(doc.paragraphs) > 9:
        p9 = doc.paragraphs[9]
        if "Era Realtor Sdn. Bhd. [E(1)2053/1]" in p9.text:
            p9.text = p9.text.replace("Era Realtor Sdn. Bhd. [E(1)2053/1]", agency_name)

    # 5. Signatures (Paragraph 16 & 17)
    signer_name = data.get("customer_signer", cust_name)
    staff_name = data.get("staff_name", "Leong Chu Ping")
    signer_date = data.get("signer_date", header_date)
    staff_date = data.get("staff_date", header_date)

    p16 = doc.paragraphs[16]
    if len(p16.runs) > 1:
        p16.runs[1].text = f" {signer_name}"
    if len(p16.runs) > 11:
        p16.runs[11].text = staff_name

    p17 = doc.paragraphs[17]
    if len(p17.runs) > 2:
        p17.runs[2].text = signer_date
    if len(p17.runs) > 12:
        p17.runs[12].text = staff_date

    # 6. Page 2 Header Date (Paragraph 19)
    p19 = doc.paragraphs[19]
    if len(p19.runs) > 16:
        p19.runs[16].text = header_date

    # 7. Table 1 (Properties Viewed)
    t1 = doc.tables[1]
    props = data.get("properties_viewed", [])
    if props:
        for idx, prop in enumerate(props):
            row_idx = idx + 1
            if row_idx < len(t1.rows):
                row = t1.rows[row_idx]
                if prop.get("no"):
                    p_no = row.cells[0].paragraphs[-1]
                    p_no.text = prop["no"]
                if prop.get("details"):
                    p_det = row.cells[1].paragraphs[0]
                    p_det.text = f"\n{prop['details']}"
                if prop.get("date"):
                    p_dt = row.cells[2].paragraphs[0]
                    p_dt.text = f"\n{prop['date']}"
                if prop.get("price"):
                    price_val = prop["price"]
                    p_pr0 = row.cells[3].paragraphs[0]
                    p_pr0.text = "\nSelling Price:"
                    if len(row.cells[3].paragraphs) > 1:
                        row.cells[3].paragraphs[1].text = price_val.replace("Selling Price:", "").strip()
                    else:
                        row.cells[3].add_paragraph(price_val.replace("Selling Price:", "").strip())
                if prop.get("description"):
                    desc = prop["description"]
                    # Split lines if formatted with bullets
                    lines = desc.split("\n")
                    for l_i, line in enumerate(lines):
                        if l_i < len(row.cells[4].paragraphs):
                            if l_i == 0:
                                row.cells[4].paragraphs[0].text = f"\n{line}"
                            else:
                                row.cells[4].paragraphs[l_i].text = line
                        else:
                            row.cells[4].add_paragraph(line)

    # 8. Table 2 (Submission of Documents)
    t2 = doc.tables[2]
    docs_sub = data.get("documents_submitted", [])
    if docs_sub and len(t2.rows) > 1:
        row = t2.rows[1]
        doc_item = docs_sub[0]
        if doc_item.get("qty"):
            row.cells[2].paragraphs[0].text = f"\n{doc_item['qty']}"

    # 9. Page 2 Bottom Date (Paragraph 26)
    if len(doc.paragraphs) > 26:
        p26 = doc.paragraphs[26]
        if len(p26.runs) > 0:
            p26.runs[0].text = header_date

    return doc

# Test with Nick data
data_nick = {
    "form_no": "0257",
    "date": "27.03.2025",
    "customer_name": "Nick",
    "no_of_pax": "2",
    "company_name": "ELPIJI (M) SDN BHD",
    "company_address": "601-A, Level 6, Tower A, Uptown 5, 5, Jalan SS21/39,\nDamansara Uptown, 47400 Petaling Jaya, Selangor D. E., Malaysia",
    "phone": "+60126761818",
    "called_in_date": "17-03-2025",
    "customer_request": "Required 2-3 acres Industrial Factory & Land with TNB Supply: 400amp",
    "requirement_summary": "Industrial Factory and Land",
    "target_location": "Temerloh",
    "remarks": "Currently Industrial Factory has 300amp",
    "customer_signer": "Nick",
    "staff_name": "Leong Chu Ping",
    "partner_agency": "Era Realtor Sdn. Bhd. [E(1)2053/1]",
    "properties_viewed": [
        {
            "no": "1.",
            "details": "2.7 acres Mentakab Industry Land For Sale LOT 1745",
            "date": "27.03.2025",
            "price": "RM 4,000,000",
            "description": "• Freehold & Industrial\nTitle\n• Connected electrical &\nwater supply and\nFencing & Gated"
        }
    ],
    "documents_submitted": [
        {
            "qty": "1 set"
        }
    ]
}

doc_out = populate_doc(data_nick)
doc_out.save("artifacts/Generated_Nick_Docx.docx")
print("Saved artifacts/Generated_Nick_Docx.docx")

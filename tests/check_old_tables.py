import docx
from xml.etree import ElementTree as ET

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
doc_o = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")

t0_c = doc_c.tables[0]
t0_o = doc_o.tables[0]

root_o0 = ET.fromstring(t0_o.rows[0].cells[0]._tc.xml)
root_o1 = ET.fromstring(t0_o.rows[0].cells[1]._tc.xml)

print("Old Cell 0 child tags:", [elem.tag.split('}')[-1] for elem in root_o0])
print("Old Cell 0 nested tables:", len(root_o0.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")))
print("Old Cell 1 child tags:", [elem.tag.split('}')[-1] for elem in root_o1])
print("Old Cell 1 nested tables:", len(root_o1.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")))

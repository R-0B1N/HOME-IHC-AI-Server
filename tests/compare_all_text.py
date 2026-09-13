import docx

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
doc_o = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")

print("=== DOC_C (CORRECT FORMAT) ALL TEXT STRINGS ===")
for i, p in enumerate(doc_c.paragraphs):
    if p.text.strip():
        print(f"P{i:02d}: {p.text}")

print("\n=== DOC_O (OLD TEMPLATE) ALL TEXT STRINGS ===")
for i, p in enumerate(doc_o.paragraphs):
    if p.text.strip():
        print(f"P{i:02d}: {p.text}")

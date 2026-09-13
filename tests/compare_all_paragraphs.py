import docx
from xml.etree import ElementTree as ET

correct_doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
old_doc = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")

print("=== PARAGRAPH BY PARAGRAPH COMPARISON (BODY) ===")
for idx in range(max(len(correct_doc.paragraphs), len(old_doc.paragraphs))):
    p_c = correct_doc.paragraphs[idx] if idx < len(correct_doc.paragraphs) else None
    p_o = old_doc.paragraphs[idx] if idx < len(old_doc.paragraphs) else None
    
    txt_c = p_c.text if p_c else "<MISSING>"
    txt_o = p_o.text if p_o else "<MISSING>"
    
    xml_c = ET.tostring(p_c._p, encoding='unicode') if p_c else ""
    xml_o = ET.tostring(p_o._p, encoding='unicode') if p_o else ""
    
    same = (txt_c == txt_o)
    print(f"\n--- Paragraph {idx} (Same text: {same}) ---")
    print(f"  CORRECT text: {txt_c!r}")
    print(f"  OLD text:     {txt_o!r}")
    if not same:
        pPr_c = p_c._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
        pPr_o = p_o._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
        print(f"  CORRECT pPr: {ET.tostring(pPr_c, encoding='unicode') if pPr_c is not None else 'None'}")
        print(f"  OLD pPr:     {ET.tostring(pPr_o, encoding='unicode') if pPr_o is not None else 'None'}")

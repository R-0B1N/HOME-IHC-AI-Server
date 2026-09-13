import docx

doc = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")
print("=== TEMPLATE IN BACKEND APP TEMPLATES ===")
c0 = doc.tables[0].rows[0].cells[0]
print("Cell 0 paragraphs count:", len(c0.paragraphs))
for i, p in enumerate(c0.paragraphs):
    print(f"c0[{i:02d}]: {p.text!r}")

c1 = doc.tables[0].rows[0].cells[1]
print("\nCell 1 paragraphs count:", len(c1.paragraphs))
for i, p in enumerate(c1.paragraphs):
    print(f"c1[{i:02d}]: {p.text!r}")

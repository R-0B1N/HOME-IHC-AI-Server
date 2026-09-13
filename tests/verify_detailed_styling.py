import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor

golden_doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
gen_doc = docx.Document("artifacts/Generated_Nick_Docx.docx")

# 1. Company Name highlight
comp_p_g = golden_doc.tables[0].rows[0].cells[0].paragraphs[4]
comp_p_gen = gen_doc.tables[0].rows[0].cells[0].paragraphs[4]
print("Golden Company Name text:", comp_p_g.text)
print("Gen Company Name text:", comp_p_gen.text)
print("Golden runs highlights:", [r.font.highlight_color for r in comp_p_g.runs])
print("Gen runs highlights:", [r.font.highlight_color for r in comp_p_gen.runs])
assert comp_p_gen.runs[1].font.highlight_color == WD_COLOR_INDEX.YELLOW

# 2. Form No red runs
p5_g = golden_doc.paragraphs[5]
p5_gen = gen_doc.paragraphs[5]
print("\nGolden P5 text:", p5_g.text)
print("Gen P5 text:", p5_gen.text)
red_runs_gen = [r for r in p5_gen.runs if r.font.color and r.font.color.rgb == RGBColor(255, 0, 0)]
print("Gen red runs count:", len(red_runs_gen))
assert len(red_runs_gen) > 0

# 3. Margins
sec_g = golden_doc.sections[0]
sec_gen = gen_doc.sections[0]
print("\nMargins Golden vs Gen:")
print(f"  Top: {sec_g.top_margin.pt} vs {sec_gen.top_margin.pt}")
print(f"  Bottom: {sec_g.bottom_margin.pt} vs {sec_gen.bottom_margin.pt}")
print(f"  Left: {sec_g.left_margin.pt} vs {sec_gen.left_margin.pt}")
print(f"  Right: {sec_g.right_margin.pt} vs {sec_gen.right_margin.pt}")
assert sec_g.top_margin.pt == sec_gen.top_margin.pt
assert sec_g.bottom_margin.pt == sec_gen.bottom_margin.pt
assert sec_g.left_margin.pt == sec_gen.left_margin.pt
assert sec_g.right_margin.pt == sec_gen.right_margin.pt

# 4. Table cell widths
print("\nTable 0 Col widths:")
for i in range(2):
    w_g = golden_doc.tables[0].rows[0].cells[i].width.pt
    w_gen = gen_doc.tables[0].rows[0].cells[i].width.pt
    print(f"  Col {i}: {w_g} vs {w_gen}")
    assert w_g == w_gen

print("\nALL STYLING, HIGHLIGHTS, COLORS, MARGINS, AND WIDTHS ASSERTIONS PASSED!")

import docx
from docx.oxml import parse_xml
from xml.etree import ElementTree as ET

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

print("==================================================")
print("SECTIONS & PAGE SETUP")
print("==================================================")
sec = doc.sections[0]
print(f"Page width: {sec.page_width.pt} pt ({sec.page_width.inches} in)")
print(f"Page height: {sec.page_height.pt} pt ({sec.page_height.inches} in)")
print(f"Top margin: {sec.top_margin.pt} pt")
print(f"Bottom margin: {sec.bottom_margin.pt} pt")
print(f"Left margin: {sec.left_margin.pt} pt")
print(f"Right margin: {sec.right_margin.pt} pt")
print(f"Header distance: {sec.header_distance.pt} pt")
print(f"Footer distance: {sec.footer_distance.pt} pt")

print("\n==================================================")
print("PARAGRAPHS & RUNS (BODY)")
print("==================================================")
for i, p in enumerate(doc.paragraphs):
    p_xml = p._p.xml
    pPr = p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
    pPr_xml = ET.tostring(pPr, encoding='unicode') if pPr is not None else ""
    print(f"\n--- Paragraph {i} ---")
    print(f"Text: {p.text!r}")
    print(f"Align: {p.alignment}")
    print(f"pPr: {pPr_xml.strip()}")
    for r_i, r in enumerate(p.runs):
        rPr = r._r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
        rPr_xml = ET.tostring(rPr, encoding='unicode') if rPr is not None else ""
        color = r.font.color.rgb if r.font.color else None
        print(f"  Run {r_i}: text={r.text!r}, font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, bold={r.bold}, italic={r.italic}, underline={r.underline}, color={color}, hl={r.font.highlight_color}")
        print(f"         rPr: {rPr_xml.strip()}")

print("\n==================================================")
print("TABLES (MAIN)")
print("==================================================")
for t_i, tbl in enumerate(doc.tables):
    tblPr = tbl._tbl.tblPr
    print(f"\n*** MAIN TABLE {t_i} ***")
    print(f"Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}")
    tblBorders = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders')
    print(f"tblBorders: {ET.tostring(tblBorders, encoding='unicode') if tblBorders is not None else 'None'}")
    tblW = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblW')
    print(f"tblW: {ET.tostring(tblW, encoding='unicode') if tblW is not None else 'None'}")

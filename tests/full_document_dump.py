import docx
from xml.etree import ElementTree as ET

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

def get_text_and_num(p):
    numPr = p._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
    num_info = ""
    if numPr is not None:
        nid = numPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        num_info = f"[numId={nid}] "
    return num_info + p.text

def dump_cell(cell, prefix=""):
    print(f"{prefix}Cell (w={cell.width.pt if cell.width else 'None'}pt):")
    for child in cell._tc:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(child, cell)
            print(f"{prefix}  P: {get_text_and_num(p)!r}")
        elif tag == 'tbl':
            tbl = docx.table.Table(child, cell)
            print(f"{prefix}  [NESTED TABLE rows={len(tbl.rows)}, cols={len(tbl.columns)}]")
            for r_i, row in enumerate(tbl.rows):
                for c_i, c in enumerate(row.cells):
                    for cp in c.paragraphs:
                        if cp.text.strip() or cp._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr') is not None:
                            print(f"{prefix}    R{r_i} C{c_i}: {get_text_and_num(cp)!r}")

print("=== BODY PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    print(f"P{i:02d}: {get_text_and_num(p)!r}")

print("\n=== MAIN TABLE 0 ===")
t0 = doc.tables[0]
dump_cell(t0.rows[0].cells[0], "T0 C0: ")
dump_cell(t0.rows[0].cells[1], "T0 C1: ")

print("\n=== MAIN TABLE 1 ===")
t1 = doc.tables[1]
for r_i, r in enumerate(t1.rows):
    print(f"Row {r_i}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")

print("\n=== MAIN TABLE 2 ===")
t2 = doc.tables[2]
for r_i, r in enumerate(t2.rows):
    print(f"Row {r_i}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")
    if r_i == 1:
        dump_cell(r.cells[1], "T2 R1 C1: ")
        dump_cell(r.cells[4], "T2 R1 C4: ")

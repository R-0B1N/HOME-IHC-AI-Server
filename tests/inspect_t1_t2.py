import docx
from xml.etree import ElementTree as ET

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

def inspect_table_full(tbl, name):
    print(f"\n=================== {name} ===================")
    for r_i, row in enumerate(tbl.rows):
        print(f"--- Row {r_i} ---")
        for c_i, cell in enumerate(row.cells):
            tcPr = cell._tc.tcPr
            tcW = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW')
            w_val = tcW.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w') if tcW is not None else None
            # get all paragraphs and nested tables
            p_list = []
            for child in cell._tc:
                tag = child.tag.split('}')[-1]
                if tag == 'p':
                    # get text and numPr
                    numPr = child.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
                    numId = ""
                    if numPr is not None:
                        numId = " [numId=" + numPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') + "]"
                    t_str = "".join([t.text for t in child.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text])
                    p_list.append(f"P({repr(t_str)}{numId})")
                elif tag == 'tbl':
                    p_list.append("NESTED_TBL")
            print(f"  Col {c_i} (w={w_val} dxa): {', '.join(p_list)}")

inspect_table_full(doc.tables[1], "TABLE 1")
inspect_table_full(doc.tables[2], "TABLE 2")

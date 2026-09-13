import sys, os
sys.path.insert(0, os.path.abspath("."))
import docx
from tests.test_complete_population import populate_doc

sample_0190 = {
    "form_no": "0190",
    "date": "27.03.2025",
    "salutation": "MR",
    "customer_name": "Lee Wan Soon",
    "no_of_pax": "2",
    "company_name": "ELPIJI (M) SDN BHD",
    "company_reg_no": "",
    "company_address": "601-A, Level 6, Tower A, Uptown 5, 5, Jalan SS21/39, Damansara Uptown, 47400 Petaling Jaya, Selangor D. E., Malaysia",
    "car_plate": "",
    "phone": "012-329 2280",
    "called_in_date": "17-03-2025",
    "referral_source": "Bentongland Website",
    "customer_request": "Required 2-3 acres Industrial Factory & Land with TNB Supply: 400amp",
    "requirement_summary": "Types of Properties: Industrial Factory and Land",
    "property_types": ["Industrial Land", "Factory"],
    "target_location": "Temerloh",
    "remarks": "Currently Industrial Factory has 300amp",
    "assigned_to": "Direct Seller",
    "partner_agency": "Chester Properties Sdn. Bhd [E(1)1321/16]",
    "customer_signer": "Lee Wan Soon on behalf of ELPIJI (M) SDN BHD",
    "staff_name": "Leong Chu Ping",
    "properties_viewed": [
        {
            "no": "1.",
            "details": "2.7 acres Mentakab Industry Land For Sale LOT 1745",
            "date": "27.03.2025",
            "price": "RM 4,000,000",
            "description": "• Freehold & Industrial Title\n• Connected electrical & water supply and Fencing & Gated"
        }
    ],
    "documents_submitted": [
        {
            "title_details": "Title Lot 1745",
            "qty": "1 set",
            "remarks": "Whatsapp Messenger"
        }
    ]
}

doc = populate_doc(sample_0190)
out_path = "artifacts/Viewing_Acknowledgement_0190_Corrected.docx"
doc.save(out_path)
print("Successfully generated 0190 corrected docx at", out_path)
print("Paragraphs count:", len(doc.paragraphs))
print("Tables count:", len(doc.tables))
print("P5 text:", repr(doc.paragraphs[5].text))
print("P9 text partner agency:", "Chester Properties" in doc.paragraphs[9].text)
print("P16 text:", repr(doc.paragraphs[16].text))
print("P17 text:", repr(doc.paragraphs[17].text))

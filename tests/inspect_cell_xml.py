import docx
from xml.etree import ElementTree as ET

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
doc_o = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")

t0_c = doc_c.tables[0]
t0_o = doc_o.tables[0]

print("=== CELL 0 XML TAGS ===")
c0_c_xml = t0_c.rows[0].cells[0]._tc.xml
c0_o_xml = t0_o.rows[0].cells[0]._tc.xml

# Let's see all sub-elements in cell 0
root_c = ET.fromstring(c0_c_xml)
print("Correct Cell 0 child tags:", [elem.tag.split('}')[-1] for elem in root_c])

# Are there nested tables (w:tbl) in cell 0 or cell 1?
nested_tbls_c0 = root_c.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")
print(f"Correct Cell 0 nested tables: {len(nested_tbls_c0)}")

root_c1 = ET.fromstring(t0_c.rows[0].cells[1]._tc.xml)
print("Correct Cell 1 child tags:", [elem.tag.split('}')[-1] for elem in root_c1])
nested_tbls_c1 = root_c1.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")
print(f"Correct Cell 1 nested tables: {len(nested_tbls_c1)}")

# Are there shapes, drawings, sdt (structured document tags), or pict?
print("Cell 0 drawings:", len(root_c.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")))
print("Cell 0 shapes/pict:", len(root_c.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pict")))
print("Cell 0 sdts:", len(root_c.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdt")))

print("Cell 1 drawings:", len(root_c1.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")))
print("Cell 1 shapes/pict:", len(root_c1.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pict")))
print("Cell 1 sdts:", len(root_c1.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdt")))

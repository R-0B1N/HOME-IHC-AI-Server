import docx

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
c0 = doc_c.tables[0].rows[0].cells[0]
print("=== CELL 0 PARAGRAPHS & RUNS ===")
for p_i, p in enumerate(c0.paragraphs):
    print(f"\nP{p_i:02d}: text={p.text!r}")
    for r_i, r in enumerate(p.runs):
        color = r.font.color.rgb if r.font.color else None
        print(f"  r{r_i}: text={r.text!r}, b={r.bold}, sz={r.font.size.pt if r.font.size else None}, hl={r.font.highlight_color}")

c1 = doc_c.tables[0].rows[0].cells[1]
print("\n=== CELL 1 PARAGRAPHS & RUNS ===")
for p_i, p in enumerate(c1.paragraphs):
    print(f"\nP{p_i:02d}: text={p.text!r}")
    for r_i, r in enumerate(p.runs):
        color = r.font.color.rgb if r.font.color else None
        print(f"  r{r_i}: text={r.text!r}, b={r.bold}, sz={r.font.size.pt if r.font.size else None}, hl={r.font.highlight_color}")

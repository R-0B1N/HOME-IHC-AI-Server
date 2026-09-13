import docx

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

t1 = doc_c.tables[1]
print("=== TABLE 1 ROW 1 ===")
for c_i, cell in enumerate(t1.rows[1].cells):
    print(f"\n--- Col {c_i} ---")
    for p_i, p in enumerate(cell.paragraphs):
        runs_info = [f"'{r.text}'(sz={r.font.size.pt if r.font.size else None}, b={r.bold})" for r in p.runs]
        print(f"  P{p_i}: {p.text!r} runs: {runs_info}")

t2 = doc_c.tables[2]
print("\n=== TABLE 2 ROW 1 ===")
for c_i, cell in enumerate(t2.rows[1].cells):
    print(f"\n--- Col {c_i} ---")
    for p_i, p in enumerate(cell.paragraphs):
        runs_info = [f"'{r.text}'(sz={r.font.size.pt if r.font.size else None}, b={r.bold})" for r in p.runs]
        print(f"  P{p_i}: {p.text!r} runs: {runs_info}")

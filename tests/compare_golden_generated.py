import docx
from xml.etree import ElementTree as ET

golden_doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
gen_doc = docx.Document("artifacts/Generated_Nick_Docx.docx")

print("=== BODY PARAGRAPHS COMPARISON ===")
diffs = 0
for idx in range(max(len(golden_doc.paragraphs), len(gen_doc.paragraphs))):
    p_g = golden_doc.paragraphs[idx] if idx < len(golden_doc.paragraphs) else None
    p_o = gen_doc.paragraphs[idx] if idx < len(gen_doc.paragraphs) else None
    t_g = p_g.text if p_g else "<MISSING>"
    t_o = p_o.text if p_o else "<MISSING>"
    if t_g != t_o:
        diffs += 1
        print(f"P{idx:02d} DIFF:")
        print(f"  GOLDEN:    {t_g!r}")
        print(f"  GENERATED: {t_o!r}")

if diffs == 0:
    print("ALL 27 BODY PARAGRAPHS MATCH EXACTLY!")

print("\n=== TABLE 0 COMPARISON ===")
t0_g = golden_doc.tables[0]
t0_o = gen_doc.tables[0]

def compare_cells(c_g, c_o, name):
    print(f"\nComparing {name}:")
    p_g_list = [p.text for p in c_g.paragraphs]
    p_o_list = [p.text for p in c_o.paragraphs]
    c_diff = 0
    for i in range(max(len(p_g_list), len(p_o_list))):
        tg = p_g_list[i] if i < len(p_g_list) else "<MISSING>"
        to = p_o_list[i] if i < len(p_o_list) else "<MISSING>"
        if tg != to:
            c_diff += 1
            print(f"  [{i:02d}] DIFF: GOLDEN={tg!r} vs GEN={to!r}")
    if c_diff == 0:
        print(f"  All {len(p_g_list)} paragraphs match exactly!")

compare_cells(t0_g.rows[0].cells[0], t0_o.rows[0].cells[0], "Table 0 Cell 0")
compare_cells(t0_g.rows[0].cells[1], t0_o.rows[0].cells[1], "Table 0 Cell 1")

print("\n=== TABLE 1 COMPARISON ===")
t1_g = golden_doc.tables[1]
t1_o = gen_doc.tables[1]
t1_diff = 0
for r_i in range(len(t1_g.rows)):
    for c_i in range(len(t1_g.columns)):
        tg = t1_g.rows[r_i].cells[c_i].text
        to = t1_o.rows[r_i].cells[c_i].text
        if tg != to:
            t1_diff += 1
            print(f"  T1 R{r_i} C{c_i} DIFF: GOLDEN={tg!r} vs GEN={to!r}")
if t1_diff == 0:
    print("All Table 1 cells match exactly!")

print("\n=== TABLE 2 COMPARISON ===")
t2_g = golden_doc.tables[2]
t2_o = gen_doc.tables[2]
t2_diff = 0
for r_i in range(len(t2_g.rows)):
    for c_i in range(len(t2_g.columns)):
        tg = t2_g.rows[r_i].cells[c_i].text
        to = t2_o.rows[r_i].cells[c_i].text
        if tg != to:
            t2_diff += 1
            print(f"  T2 R{r_i} C{c_i} DIFF: GOLDEN={tg!r} vs GEN={to!r}")
if t2_diff == 0:
    print("All Table 2 cells match exactly!")

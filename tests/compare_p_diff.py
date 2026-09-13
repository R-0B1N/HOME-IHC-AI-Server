import docx

doc_c = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")
doc_o = docx.Document("backend/app/templates/Property Acknowledgement Document.docx")

print("=== BODY PARAGRAPHS COUNT ===")
print("Correct:", len(doc_c.paragraphs))
print("Old:", len(doc_o.paragraphs))

print("\n=== BODY PARAGRAPHS COMPARISON ===")
for i in range(max(len(doc_c.paragraphs), len(doc_o.paragraphs))):
    tc = doc_c.paragraphs[i].text if i < len(doc_c.paragraphs) else "<NONE>"
    to = doc_o.paragraphs[i].text if i < len(doc_o.paragraphs) else "<NONE>"
    if tc != to:
        print(f"P{i:02d} DIFF:")
        print(f"  CORRECT: {tc!r}")
        print(f"  OLD:     {to!r}")
    else:
        print(f"P{i:02d} SAME: {tc[:60]!r}")

print("\n=== TABLE 0 CELL 0 PARAGRAPHS ===")
c0_c = doc_c.tables[0].rows[0].cells[0]
c0_o = doc_o.tables[0].rows[0].cells[0]
print(f"Cell 0 count - Correct: {len(c0_c.paragraphs)}, Old: {len(c0_o.paragraphs)}")
for i in range(max(len(c0_c.paragraphs), len(c0_o.paragraphs))):
    tc = c0_c.paragraphs[i].text if i < len(c0_c.paragraphs) else "<NONE>"
    to = c0_o.paragraphs[i].text if i < len(c0_o.paragraphs) else "<NONE>"
    print(f"  [{i:02d}] CORRECT: {tc!r}")
    print(f"       OLD:     {to!r}")

print("\n=== TABLE 0 CELL 1 PARAGRAPHS ===")
c1_c = doc_c.tables[0].rows[0].cells[1]
c1_o = doc_o.tables[0].rows[0].cells[1]
print(f"Cell 1 count - Correct: {len(c1_c.paragraphs)}, Old: {len(c1_o.paragraphs)}")
for i in range(max(len(c1_c.paragraphs), len(c1_o.paragraphs))):
    tc = c1_c.paragraphs[i].text if i < len(c1_c.paragraphs) else "<NONE>"
    to = c1_o.paragraphs[i].text if i < len(c1_o.paragraphs) else "<NONE>"
    print(f"  [{i:02d}] CORRECT: {tc!r}")
    print(f"       OLD:     {to!r}")

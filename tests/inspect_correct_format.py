import docx

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

print("=== CORRECT FORMAT DOCX - PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    runs_str = " | ".join([f"'{r.text}'(font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, b={r.bold}, c={r.font.color.rgb if r.font.color else None}, hl={r.font.highlight_color})" for r in p.runs])
    print(f"P{i:02d} [align={p.alignment}]: {p.text!r}")
    print(f"     runs: {runs_str}")

print("\n=== CORRECT FORMAT DOCX - TABLE 0 ===")
t0 = doc.tables[0]
for r_idx, row in enumerate(t0.rows):
    for c_idx, cell in enumerate(row.cells):
        print(f"\n--- T0 Row {r_idx}, Col {c_idx} (width={cell.width.pt if cell.width else None}pt) ---")
        for p_idx, p in enumerate(cell.paragraphs):
            runs_str = " | ".join([f"'{r.text}'(font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, b={r.bold}, c={r.font.color.rgb if r.font.color else None}, hl={r.font.highlight_color})" for r in p.runs])
            print(f"  P{p_idx:02d}: {p.text!r}")
            print(f"       {runs_str}")

print("\n=== CORRECT FORMAT DOCX - TABLE 1 ===")
t1 = doc.tables[1]
for r_idx, row in enumerate(t1.rows):
    for c_idx, cell in enumerate(row.cells):
        print(f"\n--- T1 Row {r_idx}, Col {c_idx} (width={cell.width.pt if cell.width else None}pt) ---")
        for p_idx, p in enumerate(cell.paragraphs):
            runs_str = " | ".join([f"'{r.text}'(font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, b={r.bold}, c={r.font.color.rgb if r.font.color else None})" for r in p.runs])
            print(f"  P{p_idx:02d}: {p.text!r}")
            print(f"       {runs_str}")

print("\n=== CORRECT FORMAT DOCX - TABLE 2 ===")
t2 = doc.tables[2]
for r_idx, row in enumerate(t2.rows):
    for c_idx, cell in enumerate(row.cells):
        print(f"\n--- T2 Row {r_idx}, Col {c_idx} (width={cell.width.pt if cell.width else None}pt) ---")
        for p_idx, p in enumerate(cell.paragraphs):
            runs_str = " | ".join([f"'{r.text}'(font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, b={r.bold}, c={r.font.color.rgb if r.font.color else None})" for r in p.runs])
            print(f"  P{p_idx:02d}: {p.text!r}")
            print(f"       {runs_str}")

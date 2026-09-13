import docx
from xml.etree import ElementTree as ET

doc = docx.Document("docs/Property Acknowledgement Document Correct Format.docx")

def find_checkboxes_in_cell(cell, location_name):
    # Find all nested tables
    root = ET.fromstring(cell._tc.xml)
    tbls = cell._tc.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')
    print(f"\n=== {location_name} (nested tables: {len(tbls)}) ===")
    for t_i, tbl in enumerate(tbls):
        for r_i, tr in enumerate(tbl.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')):
            cells = tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
            for c_i, tc in enumerate(cells):
                numPr = tc.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
                if numPr is not None:
                    nid = numPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    # Find label in next cell or within same row
                    label = ""
                    if c_i + 1 < len(cells):
                        label = "".join([t.text for t in cells[c_i + 1].findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]).strip()
                    print(f"  Nested Tbl {t_i} Row {r_i} Col {c_i}: numId={nid} -> Label: {label!r}")

find_checkboxes_in_cell(doc.tables[0].rows[0].cells[0], "Table 0 Cell 0 (Customer Details)")
find_checkboxes_in_cell(doc.tables[0].rows[0].cells[1], "Table 0 Cell 1 (Customer Requirements)")
find_checkboxes_in_cell(doc.tables[2].rows[1].cells[1], "Table 2 Row 1 Col 1 (Documents Submitted)")
find_checkboxes_in_cell(doc.tables[2].rows[1].cells[4], "Table 2 Row 1 Col 4 (Submission Remarks)")

import os
import docx
import fitz  # PyMuPDF

CORRECT_DOCX = "docs/Property Acknowledgement Document Correct Format.docx"
OLD_TEMPLATE = "backend/app/templates/Property Acknowledgement Document.docx"
DOCS_TEMPLATE = "docs/Property Acknowledgement Document.docx"
AI_PDF = "docs/Viewing_Acknowledgement_0190_Nick.pdf"

print("--- FILE EXISTENCE & SIZES ---")
for p in [CORRECT_DOCX, OLD_TEMPLATE, DOCS_TEMPLATE, AI_PDF]:
    if os.path.exists(p):
        print(f"{p}: {os.path.getsize(p)} bytes")
    else:
        print(f"{p}: NOT FOUND")

import filecmp
if os.path.exists(OLD_TEMPLATE) and os.path.exists(DOCS_TEMPLATE):
    print("OLD_TEMPLATE == DOCS_TEMPLATE?", filecmp.cmp(OLD_TEMPLATE, DOCS_TEMPLATE))

doc_correct = docx.Document(CORRECT_DOCX)
doc_old = docx.Document(OLD_TEMPLATE)

print(f"\nCorrect docx paragraphs: {len(doc_correct.paragraphs)}, tables: {len(doc_correct.tables)}, sections: {len(doc_correct.sections)}")
print(f"Old docx paragraphs: {len(doc_old.paragraphs)}, tables: {len(doc_old.tables)}, sections: {len(doc_old.sections)}")

print("\n--- SECTIONS & MARGINS ---")
for name, d in [("Correct", doc_correct), ("Old", doc_old)]:
    sec = d.sections[0]
    print(f"{name}: Page Width={sec.page_width.pt}pt, Height={sec.page_height.pt}pt, Top={sec.top_margin.pt}pt, Bottom={sec.bottom_margin.pt}pt, Left={sec.left_margin.pt}pt, Right={sec.right_margin.pt}pt")

print("\n--- PARAGRAPHS COMPARISON ---")
print("Correct doc paragraphs:")
for i, p in enumerate(doc_correct.paragraphs):
    print(f"[{i}] (align={p.alignment}, runs={len(p.runs)}): {repr(p.text)}")
    for r_idx, r in enumerate(p.runs):
        font_name = r.font.name
        font_sz = r.font.size.pt if r.font.size else None
        color = r.font.color.rgb if r.font.color else None
        print(f"    run[{r_idx}] ({font_name}, {font_sz}pt, bold={r.bold}, color={color}): {repr(r.text)}")

print("\nOld doc paragraphs:")
for i, p in enumerate(doc_old.paragraphs):
    print(f"[{i}] (align={p.alignment}, runs={len(p.runs)}): {repr(p.text)}")
    for r_idx, r in enumerate(p.runs):
        font_name = r.font.name
        font_sz = r.font.size.pt if r.font.size else None
        color = r.font.color.rgb if r.font.color else None
        print(f"    run[{r_idx}] ({font_name}, {font_sz}pt, bold={r.bold}, color={color}): {repr(r.text)}")

print("\n--- TABLES SUMMARY ---")
for t_idx in range(max(len(doc_correct.tables), len(doc_old.tables))):
    t_c = doc_correct.tables[t_idx] if t_idx < len(doc_correct.tables) else None
    t_o = doc_old.tables[t_idx] if t_idx < len(doc_old.tables) else None
    print(f"\nTable {t_idx}:")
    if t_c:
        print(f"  Correct: rows={len(t_c.rows)}, cols={len(t_c.columns)}")
    if t_o:
        print(f"  Old: rows={len(t_o.rows)}, cols={len(t_o.columns)}")

print("\n--- PDF INSPECTION ---")
if os.path.exists(AI_PDF):
    pdf = fitz.open(AI_PDF)
    print(f"PDF page count: {len(pdf)}")
    for page_num in range(len(pdf)):
        page = pdf[page_num]
        print(f"Page {page_num+1} rect: {page.rect}")
        text = page.get_text()
        print(f"Page {page_num+1} text length: {len(text)}, first 200 chars:\n{repr(text[:200])}")

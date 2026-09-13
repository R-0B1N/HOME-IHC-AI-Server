"""
Automated Test Suite for Sprint 2: Logic, Schedulers, Transcripts & Document Engine.
Tests:
1. Property Acknowledgement Document Engine (22 fields, XML tblBorders, Unicode checkboxes, yellow highlight, bold red run).
2. Branded Chatwoot PDF Transcript Generator (PyMuPDF vector PDF, metadata, bubbles, multi-page).
3. OpenXML Reporting Service (Buyer Database.xlsx and Owner Database.xlsx).
4. Accelerated Lead Nurturing Daemon & Meta 24-Hour Messaging Policy Guard.
5. Celery Beat Schedule Configuration.
6. Chatwoot Private Note Agent Commands (/acknowledgement and /transcript).
"""

import os
import io
import json
import zipfile
import unittest
from unittest.mock import patch, MagicMock
from datetime import datetime, timezone, timedelta

import docx
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX

from app.services.acknowledgement import (
    populate_acknowledgement_document,
    generate_viewing_acknowledgement,
    get_sample_acknowledgement_data,
    stabilize_table_borders,
)
from app.services.transcript import generate_conversation_transcript_pdf
from app.services.reporting import (
    create_xlsx_bytes,
    export_buyer_database_xlsx,
    export_owner_database_xlsx,
    generate_weekly_database_reports,
)
from app.worker.celery_app import celery_app
from app.worker.tasks import run_lead_nurturing_daemon, generate_and_send_weekly_reports


class TestAcknowledgementEngine(unittest.TestCase):
    """Tests the Property Viewing Acknowledgement Word and PDF engine."""

    def test_sample_data_generation_and_openxml_borders(self):
        sample_data = get_sample_acknowledgement_data()
        self.assertEqual(sample_data["form_no"], "0190")
        self.assertEqual(sample_data["customer_name"], "Lee Wan Soon")

        doc = populate_acknowledgement_document(sample_data)
        
        # Verify Table 0 exists and has XML tblBorders
        self.assertGreater(len(doc.tables), 0)
        t0 = doc.tables[0]
        self.assertIn("tblBorders", t0._tbl.tblPr.xml)
        self.assertIn('w:val="single"', t0._tbl.tblPr.xml)

    def test_22_fields_population(self):
        data = get_sample_acknowledgement_data()
        data["form_no"] = "0999"
        data["customer_name"] = "Tan Sri Dato Robert"
        data["company_name"] = "MEGA CAPITAL BERHAD"
        data["company_reg_no"] = "202301099999"
        data["phone"] = "019-888 7777"
        data["salutation"] = "MR"
        data["referral_source"] = "Mudah"
        data["partner_agency"] = "Chester Properties Sdn. Bhd [E(1)1321/16]"
        data["property_types"] = ["Industrial Land", "Factory"]
        data["target_location"] = "Karak"
        data["remarks"] = "Requires 3-phase power 600amp"

        doc = populate_acknowledgement_document(data)

        # 1. Form No in P5 (red bold run)
        p5 = doc.paragraphs[5]
        self.assertIn("0999", p5.text)
        red_runs = [r for r in p5.runs if r.font.color and r.font.color.rgb == RGBColor(255, 0, 0)]
        self.assertGreater(len(red_runs), 0)
        self.assertTrue(red_runs[0].bold)

        # 2. Header Date in P7
        p7 = doc.paragraphs[7]
        self.assertIn("27.03.2025", p7.text)

        # 3. Table 0 Cell 0: Salutation in nested table 0
        c0 = doc.tables[0].rows[0].cells[0]
        c0_nested = [docx.table.Table(node, c0) for node in c0._tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')]
        mr_num = c0_nested[0].rows[0].cells[0].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        mrs_num = c0_nested[0].rows[0].cells[2].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        self.assertEqual(mr_num, "999")
        self.assertEqual(mrs_num, "998")

        # 4. Full Name in c0.paragraphs[0]
        self.assertIn("Tan Sri Dato Robert", c0.paragraphs[0].text)

        # 5. Pax in c0.paragraphs[2]
        self.assertIn("No of pax: 2", c0.paragraphs[2].text)

        # 6. Company Name & Yellow Highlight in c0.paragraphs[4]
        self.assertIn("MEGA CAPITAL BERHAD", c0.paragraphs[4].text)
        highlighted_runs = [r for r in c0.paragraphs[4].runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertGreater(len(highlighted_runs), 0)

        # 7. Phone in c0.paragraphs[10]
        self.assertIn("019-888 7777", c0.paragraphs[10].text)

        # 8. Referral Source toggling in c0 nested table 1
        mudah_num = c0_nested[1].rows[1].cells[0].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        bentong_num = c0_nested[1].rows[2].cells[0].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        self.assertEqual(mudah_num, "999")
        self.assertEqual(bentong_num, "998")

        # 9. Table 0 Cell 1: Customer request and property types
        c1 = doc.tables[0].rows[0].cells[1]
        self.assertIn("Customer’s request:", c1.paragraphs[1].text)
        c1_nested = [docx.table.Table(node, c1) for node in c1._tc.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')]
        ind_num = c1_nested[0].rows[2].cells[0].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        fact_num = c1_nested[0].rows[2].cells[2].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        agri_num = c1_nested[0].rows[0].cells[0].paragraphs[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        self.assertEqual(ind_num, "999")
        self.assertEqual(fact_num, "999")
        self.assertEqual(agri_num, "998")
        self.assertIn("Location: Karak", c1.paragraphs[9].text)
        self.assertIn("Remarks: Requires 3-phase power 600amp", c1.paragraphs[11].text)

        # 10. Partner Agency legal override
        p9 = doc.paragraphs[9]
        self.assertIn("Chester Properties Sdn. Bhd [E(1)1321/16]", p9.text)
        self.assertNotIn("Era Realtor Sdn. Bhd.", p9.text)

        # 11. Page 2 Tables
        t1 = doc.tables[1]
        self.assertIn("2.7 acres Mentakab Industry", t1.rows[1].cells[1].text)
        self.assertIn("RM 4,000,000", t1.rows[1].cells[3].text)

        t2 = doc.tables[2]
        self.assertIn("Title", t2.rows[1].cells[1].text)
        self.assertIn("1 set", t2.rows[1].cells[2].text)

    def test_high_level_generate_acknowledgement(self):
        output_dir = "data/output/test_ack_suite"
        os.makedirs(output_dir, exist_ok=True)
        result = generate_viewing_acknowledgement(output_dir=output_dir, filename_prefix="Test_Form_0190")
        self.assertTrue(os.path.exists(result["docx_path"]))
        self.assertEqual(result["form_no"], "0190")
        self.assertEqual(result["customer_name"], "Lee Wan Soon")


class TestTranscriptEngine(unittest.TestCase):
    """Tests the PyMuPDF vector PDF conversation transcript engine."""

    @patch("app.services.chatwoot.get_conversation_messages")
    @patch("app.services.chatwoot.get_conversation_details")
    def test_transcript_pdf_generation(self, mock_details, mock_msgs):
        mock_details.return_value = {
            "id": 999,
            "status": "open",
            "meta": {
                "sender": {
                    "name": "Dato' Sri David Wong (黄先生)",
                    "phone_number": "+60123456789",
                    "email": "david@homeihc.com"
                },
                "channel": "Channel::Whatsapp"
            }
        }
        mock_msgs.return_value = [
            {
                "id": 101,
                "message_type": "incoming",
                "content": "Hello Irene, 我想了解文冬 (Bentong) 5依格的榴莲园 (Musang King durian land).",
                "created_at": 1741500000,
                "sender": {"name": "Dato' Sri David Wong"}
            },
            {
                "id": 102,
                "message_type": "outgoing",
                "content": "Hi Dato' Sri! 😊 Thank you for contacting Home IHC. We currently have an excellent 5-acre mature Musang King orchard in Karak/Bentong with stream access.",
                "created_at": 1741500060,
                "sender": {"name": "Irene Leong (Home IHC)"}
            },
            {
                "id": 103,
                "message_type": "outgoing",
                "content": "Confidential internal note",
                "created_at": 1741500090,
                "private": True,
                "sender": {"name": "Staff Note"}
            }
        ]

        # 1. Without private notes
        pdf_bytes, filename = generate_conversation_transcript_pdf(conversation_id=999, include_private_notes=False)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertIn("Transcript_Conv_999", filename)
        self.assertGreater(len(pdf_bytes), 1000)

        # 2. Open and inspect using PyMuPDF (fitz)
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        self.assertGreater(len(doc), 0)
        page1_text = doc[0].get_text()
        self.assertIn("HOME IHC", page1_text)
        self.assertIn("BentongLand.com.my", page1_text)
        self.assertIn("999", page1_text)
        self.assertNotIn("Confidential internal note", page1_text)


class TestReportingEngine(unittest.TestCase):
    """Tests the Excel OpenXML database reporting engine."""

    def test_create_xlsx_bytes(self):
        headers = ["Customer ID", "Name", "Phone", "Status"]
        rows = [
            ["C101", "Ahmad bin Abdullah", "+60191234567", "Hot Lead"],
            ["C102", "Chong Wei & Partners", "+60129876543", "Under Offer"]
        ]
        xlsx_bytes = create_xlsx_bytes(headers, rows, sheet_name="Test DB")
        self.assertGreater(len(xlsx_bytes), 500)

        # Inspect zipfile structure
        buf = io.BytesIO(xlsx_bytes)
        with zipfile.ZipFile(buf, "r") as zf:
            namelist = zf.namelist()
            self.assertIn("[Content_Types].xml", namelist)
            self.assertIn("xl/workbook.xml", namelist)
            self.assertIn("xl/worksheets/sheet1.xml", namelist)

    def test_export_buyer_and_owner_database_with_mock(self):
        mock_session = MagicMock()

        # Mock Customers
        c1 = MagicMock()
        c1.id = "60123456789"
        c1.contact_name = "Mr. Tan"
        c1.phone_number = "+60123456789"
        c1.email = "tan@example.com"
        c1.created_at = datetime.now()
        c1.updated_at = datetime.now()
        c1.metadata_json = {
            "intent": "buyer",
            "lead_temp": "hot",
            "buyer_location": "Bentong",
            "buyer_property_type": "Durian Orchard",
            "buyer_budget": "RM 1.5M - 2M",
            "bypass_ai": False
        }

        # Mock Properties
        p1 = MagicMock()
        p1.id = "prop-uuid-001"
        p1.title = "5 Acres Karak Durian Farm"
        p1.listing_status = "Available"
        p1.property_category = ["Agriculture", "Durian Land"]
        p1.property_type_sub = "Agricultural"
        p1.asking_price_myr = 1750000.0
        p1.price_per_acre_myr = 350000.0
        p1.price_per_sqft_myr = 8.03
        p1.land_area_acres = 5.0
        p1.land_area_sqft = 217800.0
        p1.state = "Pahang"
        p1.city = "Karak"
        p1.tenure_type = "Freehold"
        p1.title_status = "Geran Mukim"
        p1.agent_name = "Irene Leong"
        p1.agent_phone = "+6011-65144931"
        p1.source_url = "https://bentongland.com.my/karak-durian-farm"

        def mock_query(model):
            q = MagicMock()
            if model.__name__ == "Customer":
                q.all.return_value = [c1]
            else:
                q.all.return_value = [p1]
            return q

        mock_session.query.side_effect = mock_query

        out_dir = "data/output/test_reports_suite"
        os.makedirs(out_dir, exist_ok=True)
        reports = generate_weekly_database_reports(output_dir=out_dir, db_session=mock_session)

        self.assertTrue(os.path.exists(reports["buyer_report"]))
        self.assertTrue(os.path.exists(reports["owner_report"]))
        self.assertTrue(os.path.exists(reports["canonical_buyer"]))
        self.assertTrue(os.path.exists(reports["canonical_owner"]))


class TestLeadNurturingDaemon(unittest.TestCase):
    """Tests the lead nurturing daemon and WhatsApp 24-hour customer care policy guard."""

    @patch("app.worker.tasks.SessionLocal")
    @patch("app.worker.tasks.send_message")
    @patch("app.worker.tasks.send_private_note")
    def test_lead_nurturing_within_and_outside_24h(self, mock_note, mock_msg, mock_session_local):
        mock_db = MagicMock()
        mock_session_local.return_value = mock_db

        now = datetime.now(timezone.utc)

        # 1. Customer A: Hot Lead inactive for 20 hours (WITHIN 24h window)
        cust_a = MagicMock()
        cust_a.id = "cust_a"
        cust_a.contact_name = "Alex"
        cust_a.phone_number = "+60111111111"
        cust_a.created_at = now - timedelta(hours=20)
        cust_a.updated_at = now - timedelta(hours=20)
        cust_a.metadata_json = {
            "lead_temp": "hot",
            "conversation_id": 1001,
            "bypass_ai": False
        }

        # 2. Customer B: Hot Lead inactive for 48 hours (OUTSIDE 24h window -> Meta policy guard)
        cust_b = MagicMock()
        cust_b.id = "cust_b"
        cust_b.contact_name = "Brian"
        cust_b.phone_number = "+60122222222"
        cust_b.created_at = now - timedelta(hours=48)
        cust_b.updated_at = now - timedelta(hours=48)
        cust_b.metadata_json = {
            "lead_temp": "hot",
            "conversation_id": 1002,
            "bypass_ai": False
        }

        # 3. Customer C: Bypassed customer
        cust_c = MagicMock()
        cust_c.id = "cust_c"
        cust_c.metadata_json = {"bypass_ai": True}

        mock_db.query.return_value.all.return_value = [cust_a, cust_b, cust_c]

        result = run_lead_nurturing_daemon()
        self.assertEqual(result["status"], "success")

        # Customer A (within 24h): Automated message sent
        mock_msg.assert_called_once()
        args, kwargs = mock_msg.call_args
        self.assertEqual(args[0], 1001)
        self.assertIn("Alex", args[1])

        # Customer B (outside 24h): Private note posted instead of direct message
        mock_note.assert_called_once()
        note_args, note_kwargs = mock_note.call_args
        self.assertEqual(note_args[0], 1002)
        self.assertIn("WhatsApp 24-Hour Policy Window Closed", note_args[1])
        self.assertIn("Brian", note_args[1])


class TestCeleryBeatSchedule(unittest.TestCase):
    """Tests that Celery Beat periodic tasks are properly registered."""

    def test_beat_schedules_registered(self):
        beat_conf = celery_app.conf.beat_schedule
        self.assertIn("run-lead-nurturing-daemon-hourly", beat_conf)
        self.assertIn("generate-weekly-reports-monday", beat_conf)
        
        nurture_entry = beat_conf["run-lead-nurturing-daemon-hourly"]
        self.assertEqual(nurture_entry["task"], "app.worker.tasks.run_lead_nurturing_daemon")

        report_entry = beat_conf["generate-weekly-reports-monday"]
        self.assertEqual(report_entry["task"], "app.worker.tasks.generate_and_send_weekly_reports")


class TestAgentWebhookCommands(unittest.TestCase):
    """Tests the Chatwoot private note webhook commands for /acknowledgement and /transcript."""

    @patch("app.services.chatwoot.send_private_note")
    @patch("app.services.acknowledgement.generate_viewing_acknowledgement")
    def test_acknowledgement_private_note_command(self, mock_gen_ack, mock_note):
        from app.api.webhooks import chatwoot_webhook
        from fastapi import Request

        mock_gen_ack.return_value = {
            "docx_path": "/tmp/test.docx",
            "pdf_path": None,
            "form_no": "0190",
            "customer_name": "Test Customer"
        }

        payload = {
            "event": "message_created",
            "private": True,
            "content": "/acknowledgement",
            "conversation": {
                "id": 888,
                "meta": {
                    "sender": {
                        "name": "Test Customer",
                        "phone_number": "+60123456789"
                    }
                }
            }
        }

        raw = json.dumps(payload).encode("utf-8")
        mock_req = MagicMock(spec=Request)

        async def get_body():
            return raw

        async def get_json():
            return payload

        mock_req.body = get_body
        mock_req.json = get_json
        mock_req.headers = {}

        import asyncio
        resp = asyncio.run(chatwoot_webhook(mock_req))
        self.assertEqual(resp.get("status"), "command_executed")
        self.assertEqual(resp.get("command"), "/acknowledgement")
        mock_note.assert_called_once()
        self.assertIn("Customer Property Viewing Acknowledgement Generated", mock_note.call_args[0][1])

    @patch("app.services.chatwoot.send_private_note")
    @patch("app.services.transcript.generate_conversation_transcript_pdf")
    def test_transcript_private_note_command(self, mock_gen_trans, mock_note):
        from app.api.webhooks import chatwoot_webhook
        from fastapi import Request

        mock_gen_trans.return_value = (b"%PDF-1.4 test", "Transcript_Conv_888.pdf")

        payload = {
            "event": "message_created",
            "private": True,
            "content": "/transcript",
            "conversation": {
                "id": 888
            }
        }

        raw = json.dumps(payload).encode("utf-8")
        mock_req = MagicMock(spec=Request)

        async def get_body():
            return raw

        async def get_json():
            return payload

        mock_req.body = get_body
        mock_req.json = get_json
        mock_req.headers = {}

        import asyncio
        resp = asyncio.run(chatwoot_webhook(mock_req))
        self.assertEqual(resp.get("status"), "command_executed")
        self.assertEqual(resp.get("command"), "/transcript")
        mock_note.assert_called_once()
        self.assertIn("Conversation Transcript Generated", mock_note.call_args[0][1])

    @patch("app.services.chatwoot.send_private_note")
    def test_reset_private_note_command(self, mock_note):
        from app.api.webhooks import chatwoot_webhook
        from fastapi import Request

        payload = {
            "event": "message_created",
            "private": True,
            "content": "/reset",
            "conversation": {
                "id": 888,
                "meta": {
                    "sender": {
                        "name": "Test Customer",
                        "phone_number": "+60123456789"
                    }
                }
            }
        }

        raw = json.dumps(payload).encode("utf-8")
        mock_req = MagicMock(spec=Request)

        async def get_body():
            return raw

        async def get_json():
            return payload

        mock_req.body = get_body
        mock_req.json = get_json
        mock_req.headers = {}

        import asyncio
        resp = asyncio.run(chatwoot_webhook(mock_req))
        self.assertEqual(resp.get("status"), "command_executed")
        self.assertEqual(resp.get("command"), "/reset")
        mock_note.assert_called_once()
        self.assertIn("Session & AI State Reset Complete", mock_note.call_args[0][1])


if __name__ == "__main__":
    unittest.main()

